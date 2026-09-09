"""Generate the study's vertically oriented, styled Excel workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CSV_DELIMITER = ";"


def _read_csv(csv_path: str | Path) -> tuple[list[str], list[dict]]:
    with Path(csv_path).open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=CSV_DELIMITER)
        return list(reader.fieldnames or []), list(reader)


def _safe_excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def generate_styled_excel(
    *,
    rows: list[dict] | None = None,
    headers: list[str] | None = None,
    excel_path: str | Path = "wyniki_badania_sformatowane.xlsx",
    csv_path: str | Path | None = None,
) -> bool:
    """Write variables as rows and participants as columns."""
    if rows is None or headers is None:
        if csv_path is None or not Path(csv_path).is_file() or Path(csv_path).stat().st_size == 0:
            return False
        headers, rows = _read_csv(csv_path)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Wyniki Badania"
    worksheet.append(
        ["Zmienna"]
        + [str(row.get("participantId") or f"Uczestnik_{index}") for index, row in enumerate(rows, 1)]
    )
    for header in headers:
        worksheet.append([header] + [_safe_excel_value(row.get(header, "")) for row in rows])

    section_fills = {
        "id": PatternFill("solid", fgColor="EAEAEA"),
        "demo": PatternFill("solid", fgColor="E6F2FF"),
        "asrs": PatternFill("solid", fgColor="E6F9E6"),
        "zwl": PatternFill("solid", fgColor="F2E6FF"),
        "trial": PatternFill("solid", fgColor="FFE6D9"),
    }
    even_fill = PatternFill("solid", fgColor="FAFAFA")
    odd_fill = PatternFill("solid", fgColor="FFFFFF")
    header_fill = PatternFill("solid", fgColor="7C3AED")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )
    header_border = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="medium", color="5D2EC0"),
        bottom=Side(style="medium", color="5D2EC0"),
    )

    worksheet.row_dimensions[1].height = 28
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border

    demographic = {"age", "gender", "education", "adhdDiagnosis", "adhdMedication"}
    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 20
        variable_cell = worksheet.cell(row_index, 1)
        variable = str(variable_cell.value)
        if variable in {"participantId", "timestamp"}:
            variable_cell.fill = section_fills["id"]
        elif variable in demographic:
            variable_cell.fill = section_fills["demo"]
        elif variable.startswith("asrs_"):
            variable_cell.fill = section_fills["asrs"]
        elif variable.startswith("zwl_") or variable == "zwlekanie_total":
            variable_cell.fill = section_fills["zwl"]
        elif variable.startswith("trial_"):
            variable_cell.fill = section_fills["trial"]
        variable_cell.font = Font(name="Segoe UI", size=10, bold=True)
        variable_cell.alignment = Alignment(horizontal="left", vertical="center")
        variable_cell.border = thin_border

        numeric = (
            variable == "age"
            or variable.startswith(("asrs_", "zwl_", "trial_"))
            or variable == "zwlekanie_total"
        )
        for column_index in range(2, worksheet.max_column + 1):
            cell = worksheet.cell(row_index, column_index)
            cell.font = Font(name="Segoe UI", size=10)
            cell.fill = even_fill if column_index % 2 == 0 else odd_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="right" if numeric else "center", vertical="center"
            )
            if numeric and isinstance(cell.value, str):
                try:
                    cell.value = float(cell.value) if "." in cell.value else int(cell.value)
                except ValueError:
                    pass

    worksheet.freeze_panes = "B2"
    for column in worksheet.columns:
        letter = get_column_letter(column[0].column)
        max_length = max((len(str(cell.value or "")) for cell in column), default=0)
        worksheet.column_dimensions[letter].width = max(max_length + 4, 12)

    Path(excel_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(excel_path)
    return True


if __name__ == "__main__":
    generate_styled_excel(csv_path="wyniki_badania.csv")

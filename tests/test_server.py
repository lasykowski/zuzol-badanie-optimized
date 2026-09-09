import csv
import io
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import server


def valid_payload(participant_id="P_20260909_001"):
    asrs_answers = {f"asrs_{item}": 0 for item in range(1, 19)}
    raw = {f"zwl_{item}": 3 for item in range(1, 41)}
    scored = {f"zwl_{item}_scored": 3 for item in range(1, 41)}
    trials = []
    targets = [5_000, 15_000, 30_000, 45_000] * 2
    for number, target in enumerate(targets, 1):
        reproduced = 0 if number == 1 else target if number == 2 else target - 100
        error = reproduced - target
        trials.append(
            {
                "trialNumber": number,
                "order": number,
                "targetMs": target,
                "actualStimulusMs": 0 if number == 1 else target,
                "reproducedMs": reproduced,
                "errorMs": error,
                "absoluteErrorMs": abs(error),
                "relativeErrorPct": round(error / target * 100, 2),
                "ratio": round(reproduced / target, 4),
                "tabHiddenDuringTrial": number == 8,
            }
        )
    return {
        "participantId": participant_id,
        "timestamp": "2026-09-09T20:00:00.000Z",
        "demographic": {
            "age": 30,
            "gender": "kobieta",
            "education": "wyzsze_magister",
            "adhdDiagnosis": "nie",
            "adhdMedication": "nie",
        },
        "asrs": {"answers": asrs_answers, "partA": 0, "partB": 0, "total": 0},
        "zwlekanie": {"rawAnswers": raw, "scoredAnswers": scored, "total": 120},
        "experiment": trials,
    }


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "DATA_DIR", tmp_path)
    monkeypatch.setattr(server, "DB_FILE", tmp_path / "wyniki_badania.sqlite3")
    monkeypatch.setattr(server, "LEGACY_CSV_FILE", tmp_path / "wyniki_badania.csv")
    server.initialize_database()
    with TestClient(server.app) as test_client:
        yield test_client


def test_valid_save_preserves_zero_and_negative_values(client):
    response = client.post("/api/save-results", json=valid_payload())
    assert response.status_code == 200
    assert response.json()["status"] == "saved"

    with sqlite3.connect(server.DB_FILE) as connection:
        stored = json.loads(connection.execute("SELECT payload_json FROM results").fetchone()[0])
    assert stored["asrs"]["total"] == 0
    assert stored["experiment"][0]["actualStimulusMs"] == 0
    assert stored["experiment"][0]["reproducedMs"] == 0
    assert stored["experiment"][0]["errorMs"] == -5000
    assert stored["experiment"][1]["errorMs"] == 0


@pytest.mark.parametrize(
    "mutate",
    [
        lambda data: data.update(participantId="../../server.py"),
        lambda data: data["demographic"].update(age=17),
        lambda data: data["asrs"].update(total=1),
        lambda data: data["experiment"].pop(),
        lambda data: data["experiment"][0].pop("actualStimulusMs"),
        lambda data: data["experiment"][0].pop("tabHiddenDuringTrial"),
        lambda data: data["experiment"][0].pop("order"),
    ],
)
def test_invalid_payloads_are_rejected(client, mutate):
    payload = valid_payload()
    mutate(payload)
    assert client.post("/api/save-results", json=payload).status_code == 422


def test_oversize_body_is_rejected(client):
    response = client.post(
        "/api/save-results",
        content=b"x" * (server.MAX_BODY_BYTES + 1),
        headers={"Content-Type": "application/json"},
    )
    assert response.status_code == 413
    assert response.headers["cache-control"] == "no-store"


def test_duplicate_is_idempotent_and_does_not_overwrite(client):
    original = valid_payload()
    changed = deepcopy(original)
    changed["demographic"]["age"] = 31
    assert client.post("/api/save-results", json=original).json()["status"] == "saved"
    duplicate = client.post("/api/save-results", json=changed)
    assert duplicate.status_code == 200
    assert duplicate.json()["status"] == "duplicate"

    with sqlite3.connect(server.DB_FILE) as connection:
        rows = connection.execute("SELECT payload_json FROM results").fetchall()
    assert len(rows) == 1
    assert json.loads(rows[0][0])["demographic"]["age"] == 30


def test_concurrent_writes(client):
    payloads = [valid_payload(f"CONCURRENT_{index:02}") for index in range(16)]

    def save(payload):
        with TestClient(server.app) as thread_client:
            return thread_client.post("/api/save-results", json=payload)

    with ThreadPoolExecutor(max_workers=8) as executor:
        responses = list(executor.map(save, payloads))
    assert all(response.status_code == 200 for response in responses)
    assert all(response.json()["status"] == "saved" for response in responses)
    with sqlite3.connect(server.DB_FILE) as connection:
        assert connection.execute("SELECT COUNT(*) FROM results").fetchone()[0] == 16


def test_static_whitelist_and_headers(client):
    index = client.get("/")
    assert index.status_code == 200
    assert index.headers["x-content-type-options"] == "nosniff"
    assert index.headers["cache-control"] == "no-cache"
    styles = client.get("/css/styles.css")
    assert styles.status_code == 200
    assert styles.headers["cache-control"].startswith("public")
    assert client.get("/js/app.js").status_code == 200
    for path in ("/server.py", "/format_excel.py", "/requirements.txt", "/data", "/data/file"):
        assert client.get(path).status_code == 404
    assert "access-control-allow-origin" not in index.headers


def test_csv_and_excel_exports_are_complete_and_safe(client):
    assert client.post("/api/save-results", json=valid_payload()).status_code == 200
    with sqlite3.connect(server.DB_FILE) as connection:
        connection.execute(
            "INSERT INTO legacy_results(row_json) VALUES (?)",
            (json.dumps({"participantId": "=FORMULA", "legacyNote": "@command"}),),
        )

    csv_response = client.get("/api/download/csv")
    assert csv_response.status_code == 200
    assert csv_response.content.startswith(b"\xef\xbb\xbf")
    decoded = csv_response.content.decode("utf-8-sig")
    parsed = list(csv.DictReader(io.StringIO(decoded), delimiter=";"))
    assert "trial_1_actualStimulusMs" in parsed[0]
    assert "trial_1_order" in parsed[0]
    assert "trial_1_tabHiddenDuringTrial" in parsed[0]
    assert "legacyNote" in parsed[0]
    assert parsed[0]["participantId"] == "'=FORMULA"
    assert parsed[0]["legacyNote"] == "'@command"
    assert parsed[1]["trial_1_errorMs"] == "-5000"
    assert csv_response.headers["cache-control"] == "no-store"

    excel_response = client.get("/api/download/excel")
    assert excel_response.status_code == 200
    workbook = load_workbook(io.BytesIO(excel_response.content))
    worksheet = workbook["Wyniki Badania"]
    variables = [worksheet.cell(row, 1).value for row in range(2, worksheet.max_row + 1)]
    assert "trial_1_actualStimulusMs" in variables
    assert "trial_1_tabHiddenDuringTrial" in variables
    assert worksheet.freeze_panes == "B2"


def test_health_checks_read_write_database(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok", "database": "read-write"}
    assert response.headers["cache-control"] == "no-store"


def test_one_time_legacy_csv_migration(tmp_path, monkeypatch):
    legacy = tmp_path / "wyniki_badania.csv"
    legacy.write_text(
        "\ufeffparticipantId;trial_1_errorMs;custom\nOLD_1;-5000;wartosc\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(server, "DATA_DIR", tmp_path)
    monkeypatch.setattr(server, "DB_FILE", tmp_path / "wyniki_badania.sqlite3")
    monkeypatch.setattr(server, "LEGACY_CSV_FILE", legacy)
    server.initialize_database()
    server.initialize_database()
    with sqlite3.connect(server.DB_FILE) as connection:
        assert connection.execute("SELECT COUNT(*) FROM legacy_results").fetchone()[0] == 1
    with TestClient(server.app) as test_client:
        exported = test_client.get("/api/download/csv").content.decode("utf-8-sig")
    row = list(csv.DictReader(io.StringIO(exported), delimiter=";"))[0]
    assert row["trial_1_errorMs"] == "-5000"
